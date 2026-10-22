import csv
import io
import logging
import re

from fastapi import APIRouter, Depends, FastAPI, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import ingest, rag, vectors
from .auth import current_user
from .auth import router as auth_router
from .config import settings
from .db import Anchor, CrossRef, Document, Message, Statistic, User, _uuid, get_db, init_db
from .storage import storage

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")

app = FastAPI(title="PDF Insight API", docs_url="/api/docs", openapi_url="/api/openapi.json")
if settings.cors_origins:
    app.add_middleware(CORSMiddleware, allow_origins=settings.cors_origins, allow_credentials=True,
                       allow_methods=["*"], allow_headers=["*"])
app.include_router(auth_router)
api = APIRouter(prefix="/api", dependencies=[Depends(current_user)])


@app.on_event("startup")
def _startup():
    init_db()
    vectors.client()
    ingest.resume_pending()


@app.get("/api/health")
def health():
    return {"ok": True}


def _doc_out(d: Document) -> dict:
    return {"id": d.id, "filename": d.filename, "size_bytes": d.size_bytes, "page_count": d.page_count,
            "status": d.status, "stage": d.stage, "progress": round(d.progress, 3), "error": d.error,
            "language": d.language, "ocr_pages": d.ocr_pages, "created_at": d.created_at.isoformat()}


def _owned(db: Session, user: User, document_id: str) -> Document:
    doc = db.get(Document, document_id)
    if not doc or doc.user_id != user.id:
        raise HTTPException(404, "Document not found")
    return doc


@api.get("/documents")
def list_documents(user: User = Depends(current_user), db: Session = Depends(get_db)):
    docs = db.scalars(select(Document).where(Document.user_id == user.id).order_by(Document.created_at.desc())).all()
    return [_doc_out(d) for d in docs]


@api.post("/documents")
async def upload(files: list[UploadFile], user: User = Depends(current_user), db: Session = Depends(get_db)):
    limit = settings.max_upload_mb * 1024 * 1024
    created = []
    for f in files:
        data = await f.read(limit + 1)
        if len(data) > limit:
            raise HTTPException(413, f"{f.filename}: larger than {settings.max_upload_mb} MB")
        if not data.startswith(b"%PDF"):
            raise HTTPException(400, f"{f.filename}: not a PDF file")
        name = re.sub(r"[\x00-\x1f/\\]", "_", (f.filename or "document.pdf"))[:255]
        doc_id = _uuid()
        key = f"{user.id}/{doc_id}.pdf"
        storage.put(key, data)
        doc = Document(id=doc_id, user_id=user.id, filename=name, storage_key=key, size_bytes=len(data), stage="Queued")
        db.add(doc)
        db.commit()
        ingest.submit(doc.id)
        created.append(_doc_out(doc))
    return created


@api.get("/documents/{document_id}")
def get_document(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    return {**_doc_out(doc), "pages": doc.pages}


@api.post("/documents/{document_id}/reprocess")
def reprocess(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    if doc.status in ("queued", "processing"):
        raise HTTPException(409, "Already processing")
    doc.status, doc.stage, doc.progress, doc.error = "queued", "Queued", 0.0, ""
    db.commit()
    ingest.submit(doc.id)
    return _doc_out(doc)


@api.delete("/documents/{document_id}")
def delete_document(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    if doc.status == "processing":
        raise HTTPException(409, "Document is still processing")
    vectors.delete_document(doc.id)
    storage.delete(doc.storage_key)
    db.delete(doc)
    db.commit()
    return {"ok": True}


@api.get("/documents/{document_id}/file")
def document_file(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    return StreamingResponse(storage.stream(doc.storage_key), media_type="application/pdf", headers={
        "Cache-Control": "private, max-age=3600", "X-Content-Type-Options": "nosniff",
        "Content-Disposition": "inline"})


@api.get("/documents/{document_id}/links")
def links(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Cross-reference anchors + in-text mentions, used to draw clickable links over the PDF."""
    doc = _owned(db, user, document_id)
    anchors = db.scalars(select(Anchor).where(Anchor.document_id == doc.id)).all()
    xrefs = db.scalars(select(CrossRef).where(CrossRef.document_id == doc.id)).all()
    return {
        "anchors": [{"key": a.key, "label": a.label, "page": a.page, "rects": a.rects} for a in anchors],
        "xrefs": [{"key": x.key, "text": x.text, "page": x.page, "rects": x.rects} for x in xrefs],
    }


def _stats(db: Session, doc: Document) -> list[Statistic]:
    return db.scalars(select(Statistic).where(Statistic.document_id == doc.id)
                      .order_by(Statistic.page, Statistic.source, Statistic.id)).all()


@api.get("/documents/{document_id}/statistics")
def statistics(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    return [{"id": s.id, "page": s.page, "value": s.value, "numeric_value": s.numeric_value, "unit": s.unit,
             "metric": s.metric, "category": s.category, "context": s.context, "source": s.source, "rects": s.rects}
            for s in _stats(db, doc)]


_EXPORT_COLS = ["Value", "Metric", "Category", "Unit", "Context", "Page", "Source"]


def _safe_cell(v):
    # Neutralise spreadsheet formula injection from document text.
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


@api.get("/documents/{document_id}/statistics/export")
def export_statistics(document_id: str, format: str = "xlsx", user: User = Depends(current_user),
                      db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    rows = [[s.value, s.metric, s.category, s.unit, s.context, s.page, s.source] for s in _stats(db, doc)]
    base = re.sub(r"[^\w.-]+", "_", doc.filename.rsplit(".", 1)[0])[:80] or "document"
    if format == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")  # BOM so Excel opens Arabic CSV as UTF-8
        w = csv.writer(buf)
        w.writerow(_EXPORT_COLS)
        w.writerows([[_safe_cell(c) for c in r] for r in rows])
        return StreamingResponse(iter([buf.getvalue().encode("utf-8")]), media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": f'attachment; filename="{base}_statistics.csv"'})
    if format != "xlsx":
        raise HTTPException(400, "format must be csv or xlsx")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Statistics"
    ws.append(_EXPORT_COLS)
    for r in rows:
        ws.append([_safe_cell(c) for c in r])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
    for col, width in zip("ABCDEFG", (18, 40, 14, 10, 90, 8, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
    if doc.language == "ar":
        ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return StreamingResponse(iter([out.getvalue()]),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{base}_statistics.xlsx"'})


def _msg_out(m: Message) -> dict:
    return {"id": m.id, "role": m.role, "content": m.content, "payload": m.payload, "created_at": m.created_at.isoformat()}


@api.get("/documents/{document_id}/messages")
def messages(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    msgs = db.scalars(select(Message).where(Message.document_id == doc.id, Message.user_id == user.id)
                      .order_by(Message.created_at)).all()
    return [_msg_out(m) for m in msgs]


@api.delete("/documents/{document_id}/messages")
def clear_messages(document_id: str, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    for m in db.scalars(select(Message).where(Message.document_id == doc.id, Message.user_id == user.id)).all():
        db.delete(m)
    db.commit()
    return {"ok": True}


class ChatIn(BaseModel):
    question: str = Field(min_length=1, max_length=4000)


@api.post("/documents/{document_id}/chat")
def chat(document_id: str, body: ChatIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    doc = _owned(db, user, document_id)
    if doc.status != "ready":
        raise HTTPException(409, "Document is not ready yet")
    result = rag.answer(db, doc, user.id, body.question.strip())
    db.add(Message(document_id=doc.id, user_id=user.id, role="user", content=body.question.strip()))
    db.commit()
    reply = Message(document_id=doc.id, user_id=user.id, role="assistant",
                    content="\n".join(s["text"] for s in result["statements"]), payload=result)
    db.add(reply)
    db.commit()
    return _msg_out(reply)


app.include_router(api)
